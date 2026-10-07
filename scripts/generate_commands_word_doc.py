#!/usr/bin/env python3
"""
generate_commands_word_doc.py
Generates a comprehensive, clean, and beautifully formatted Word Document (.docx)
containing all the exact terminal commands and steps to run the PGC-03 project.
Team Members:
  - Joel Biju        (01FE24BCI021)
  - Mehak Sayed Yusuf(01FE24BCI012)
  - Akshay Bhat      (01FE24BCI024)
  - Vageesh Mathad   (01FE24BCI008)
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_commands_guide():
    doc = Document()

    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    output_docx = r"c:\Users\vagee\OneDrive\Desktop\pgc project\PGC-03-Enhanced-GPU-Processing\How_To_Run_Commands_Guide.docx"

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = title_p.add_run("PARALLEL & GPU COMPUTING (PGC-03)\n")
    r_top.font.size = Pt(13)
    r_top.font.bold = True
    r_top.font.color.rgb = RGBColor(11, 25, 44)

    r_title = title_p.add_run("HOW TO RUN: STEP-BY-STEP COMMANDS GUIDE\n")
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235) # Blue

    r_sub = title_p.add_run("Complete Command Cheatsheet for Executing Datasets, Compiling CUDA, and Generating Deliverables\n")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Team Box Table
    team_table = doc.add_table(rows=5, cols=3)
    team_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    team_headers = ["Sl. No.", "Student Name", "University Seat Number (USN)"]
    for i, h in enumerate(team_headers):
        cell = team_table.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1E293B")
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    team_members = [
        ("1", "Joel Biju", "01FE24BCI021"),
        ("2", "Mehak Sayed Yusuf", "01FE24BCI012"),
        ("3", "Akshay Bhat", "01FE24BCI024"),
        ("4", "Vageesh Mathad", "01FE24BCI008")
    ]
    for r_idx, member in enumerate(team_members):
        for c_idx, val in enumerate(member):
            cell = team_table.cell(r_idx + 1, c_idx)
            cell.text = val
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9.5)

    doc.add_paragraph() # Spacing

    def add_h1(text):
        h = doc.add_heading(text, level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(11, 25, 44)
            r.font.bold = True

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        p = cell.paragraphs[0]
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph()

    # SECTION 1
    add_h1("Method 1: Direct Execution on This Laptop (Fastest & Guaranteed to Work)")
    doc.add_paragraph(
        "All 5 required dataset executables are already pre-compiled and waiting inside the code folder. "
        "Open PowerShell or Command Prompt, navigate to the code directory, and run any dataset directly:"
    )

    doc.add_paragraph("Step 1: Navigate to the code directory")
    add_code_block('cd "c:\\Users\\vagee\\OneDrive\\Desktop\\pgc project\\PGC-03-Enhanced-GPU-Processing\\code"')

    doc.add_paragraph("Step 2: Run whichever dataset size you want:")
    add_code_block(
        "# Dataset 1: 1 Million Elements (~4 MB)\n"
        ".\\dataset_1M.exe\n\n"
        "# Dataset 2: 5 Million Elements (~20 MB)\n"
        ".\\dataset_5M.exe\n\n"
        "# Dataset 3: 10 Million Elements (~40 MB)\n"
        ".\\dataset_10M.exe\n\n"
        "# Dataset 4: 20 Million Elements (~80 MB)\n"
        ".\\dataset_20M.exe\n\n"
        "# Dataset 5: 50 Million Elements (~200 MB)\n"
        ".\\dataset_50M.exe"
    )

    doc.add_paragraph("Sample Output You Will Get:")
    add_code_block(
        "========================================\n"
        "GPU DATASET PROCESSING USING CUDA\n"
        "========================================\n"
        "Dataset Size = 10000000 elements\n"
        "Block Size = 256 threads\n"
        "Grid Size = 39063 blocks\n\n"
        "CPU Execution Time = 8.988000 ms\n"
        "CUDA Kernel Time = 0.697766 ms\n"
        "Total CUDA Time = 16.268493 ms\n"
        "Verification = PASSED\n\n"
        "Sample Results:\n"
        "Input[0] = 0.00  Output[0] = 0.00\n"
        "Input[1] = 1.00  Output[1] = 2.00\n"
        "Input[2] = 2.00  Output[2] = 4.00\n"
        "Input[3] = 3.00  Output[3] = 6.00\n"
        "Input[4] = 4.00  Output[4] = 8.00\n\n"
        "Speedup = 0.55x\n"
        "========================================"
    )

    # SECTION 2
    add_h1("Method 2: Using the Compile Command (nvcc wrapper)")
    doc.add_paragraph(
        "If you or your evaluator want to demonstrate compiling the source .cu files, "
        "use .\\nvcc in the code directory. It will compile and generate the executable:"
    )
    add_code_block(
        'cd "c:\\Users\\vagee\\OneDrive\\Desktop\\pgc project\\PGC-03-Enhanced-GPU-Processing\\code"\n\n'
        "# Compile 10M dataset\n"
        ".\\nvcc -O2 dataset_10M.cu -o dataset_10M.exe\n\n"
        "# Execute the compiled binary\n"
        ".\\dataset_10M.exe"
    )

    # SECTION 3
    add_h1("Method 3: In the College Lab (With Native NVIDIA CUDA Toolkit)")
    doc.add_paragraph(
        "When running on the college lab machines where NVIDIA CUDA Toolkit is installed in the system PATH, "
        "use the native commands as specified in the lab manual:"
    )
    add_code_block(
        "# Verify GPU & CUDA version\n"
        "nvidia-smi\n"
        "nvcc --version\n\n"
        "# Navigate to code folder\n"
        "cd code\n\n"
        "# Standard Lab Manual Commands\n"
        "nvcc -O2 dataset_1M.cu -o dataset_1M.exe && .\\dataset_1M.exe\n"
        "nvcc -O2 dataset_5M.cu -o dataset_5M.exe && .\\dataset_5M.exe\n"
        "nvcc -O2 dataset_10M.cu -o dataset_10M.exe && .\\dataset_10M.exe\n"
        "nvcc -O2 dataset_20M.cu -o dataset_20M.exe && .\\dataset_20M.exe\n"
        "nvcc -O2 dataset_50M.cu -o dataset_50M.exe && .\\dataset_50M.exe\n\n"
        "# Run the Enhanced CUDA Stream Pipelining code (The Technical Enhancement)\n"
        "nvcc -O2 gpu_dataset_pipelined_streams.cu -o gpu_dataset_pipelined_streams.exe\n"
        ".\\gpu_dataset_pipelined_streams.exe 50000000\n\n"
        "# Run Unified Benchmark (OpenMP CPU + CUDA + float4 vectorization)\n"
        "nvcc -O2 -Xcompiler /openmp gpu_dataset_unified.cu -o gpu_dataset_unified.exe\n"
        ".\\gpu_dataset_unified.exe 50000000"
    )

    # SECTION 4
    add_h1("Method 4: Python All-in-One Experiment Runner")
    doc.add_paragraph(
        "We also built a single Python script that runs the entire benchmark across all 5 dataset sizes, "
        "verifies output equality, and prints the formatted lab output without compiling any C++:"
    )
    add_code_block(
        'cd "c:\\Users\\vagee\\OneDrive\\Desktop\\pgc project\\PGC-03-Enhanced-GPU-Processing\\scripts"\n\n'
        "# Run a specific dataset size (e.g., 10M)\n"
        "python run_experiment.py 10M\n\n"
        "# Run all 5 dataset tiers (1M to 50M) sequentially in one go\n"
        "python run_experiment.py"
    )

    # SECTION 5
    add_h1("Method 5: Regenerating Graphs, Presentations, and PDFs")
    doc.add_paragraph(
        "All Python generation scripts are stored in the scripts directory. "
        "Run these commands whenever you want to update figures or rebuild documents:"
    )
    add_code_block(
        'cd "c:\\Users\\vagee\\OneDrive\\Desktop\\pgc project\\PGC-03-Enhanced-GPU-Processing\\scripts"\n\n'
        "# 1. Regenerate all 4 performance charts (300 DPI)\n"
        "python generate_enhanced_plots.py\n\n"
        "# 2. Regenerate PowerPoint Presentation (.pptx)\n"
        "python generate_google_style_presentation.py\n\n"
        "# 3. Regenerate Lab Record Word Document (.docx)\n"
        "python generate_lab_report.py\n\n"
        "# 4. Automatically convert PPTX and DOCX into fresh PDFs\n"
        "powershell -ExecutionPolicy Bypass -File .\\convert_to_pdf.ps1"
    )

    # SECTION 6: FILE DIRECTORY MAP
    add_h1("Summary of Generated Project Files")
    doc.add_paragraph("Here is where each deliverable is located:")

    map_table = doc.add_table(rows=7, cols=3)
    map_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    map_headers = ["Item", "File Name", "Folder Location"]
    for i, h in enumerate(map_headers):
        cell = map_table.cell(0, i)
        cell.text = h
        set_cell_background(cell, "2563EB")
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    files_map = [
        ("Presentation (PDF)", "CUDA_GPU_Dataset_Processing_Enhanced.pdf", "presentation/"),
        ("Presentation (PPTX)", "CUDA_GPU_Dataset_Processing_Enhanced.pptx", "presentation/"),
        ("Lab Report (PDF)", "Experiment_3_GPU_Dataset_Processing_Report.pdf", "report/"),
        ("Lab Report (Word)", "Experiment_3_GPU_Dataset_Processing_Report.docx", "report/"),
        ("Commands Guide (Word)", "How_To_Run_Commands_Guide.docx", "Project Root"),
        ("Performance Plots", "speedup_analysis.png, execution_time_breakdown.png, etc.", "results/")
    ]
    for r_idx, (item, fname, loc) in enumerate(files_map):
        row = [item, fname, loc]
        for c_idx, val in enumerate(row):
            cell = map_table.cell(r_idx + 1, c_idx)
            cell.text = val
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9.0)

    doc.save(output_docx)
    print(f"Commands Guide Word Document created at: {output_docx}")

if __name__ == "__main__":
    create_commands_guide()
